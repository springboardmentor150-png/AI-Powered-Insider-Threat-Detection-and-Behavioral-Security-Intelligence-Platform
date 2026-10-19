import io
import csv
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any
from fastapi import APIRouter, Depends, Response, HTTPException, Query, status
from sqlalchemy.orm import Session
from sqlalchemy import or_

from app.database import get_db
from app.models import Employee, Alert, Incident, RiskScoreHistory
from app.schemas import ReportSummaryResponse, ThreatCategoryCount
from app.security import get_current_user, require_role
from app.mongodb import mongo_service
from app.risk_scoring import calculate_risk_score
from app.compliance import get_compliance_metrics

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

router = APIRouter(prefix="/reports", tags=["Reports"])


@router.get("/summary", response_model=ReportSummaryResponse)
def get_report_summary(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    total_employees = db.query(Employee).count()
    high_risk_count = db.query(Employee).filter(Employee.risk_score >= 60).count()
    total_alerts = db.query(Alert).count()
    critical_alerts = db.query(Alert).filter(Alert.severity == "CRITICAL").count()
    total_incidents = db.query(Incident).count()
    resolved_incidents = db.query(Incident).filter(Incident.status.in_(["RESOLVED", "CLOSED"])).count()

    # Threat categories
    categories = ["DATA_EXFILTRATION", "PRIVILEGE_ESCALATION", "OFF_HOURS_ACCESS", "MASS_DOWNLOAD", "UNUSUAL_LOGIN"]
    breakdown = []
    for cat in categories:
        cnt = db.query(Alert).filter(Alert.category == cat).count()
        pct = round((cnt / total_alerts) * 100, 1) if total_alerts > 0 else 0.0
        breakdown.append(ThreatCategoryCount(category=cat.replace("_", " ").title(), count=cnt, percentage=pct))

    highest_risk = db.query(Employee).order_by(Employee.risk_score.desc()).limit(5).all()

    return ReportSummaryResponse(
        generated_at=datetime.now(timezone.utc),
        total_monitored_employees=total_employees,
        high_risk_employee_count=high_risk_count,
        total_alerts_period=total_alerts,
        critical_alerts_period=critical_alerts,
        total_incidents_period=total_incidents,
        resolved_incidents_period=resolved_incidents,
        top_threat_categories=breakdown,
        highest_risk_employees=highest_risk,
    )


@router.get("/export-csv")
def export_reports_csv(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    output = io.StringIO()
    writer = csv.writer(output)

    # Header
    writer.writerow([
        "Employee ID", "Name", "Department", "Designation",
        "Risk Score", "Risk Level", "Status", "Open Alerts Count", "Incidents Count"
    ])

    employees = db.query(Employee).order_by(Employee.risk_score.desc()).all()
    for emp in employees:
        open_alerts = db.query(Alert).filter(Alert.employee_id == emp.id, Alert.status == "OPEN").count()
        incidents = db.query(Incident).filter(Incident.employee_id == emp.id).count()
        writer.writerow([
            emp.employee_id, emp.name, emp.department or "", emp.designation or "",
            emp.risk_score, emp.risk_level, emp.status, open_alerts, incidents
        ])

    csv_data = output.getvalue()
    filename = f"ITBIS_Threat_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return Response(
        content=csv_data,
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/export-excel")
def export_reports_excel(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Generates genuine multi-sheet Office OpenXML .xlsx workbook."""
    wb = openpyxl.Workbook()

    # Styling helpers
    header_fill = PatternFill(start_color="0B1626", end_color="0B1626", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="00E5FF")
    title_font = Font(name="Arial", size=14, bold=True, color="00A8FF")
    data_font = Font(name="Arial", size=9)
    border_thin = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0")
    )

    def style_header_row(ws, row_num):
        for cell in ws[row_num]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def auto_fit_columns(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    # 1. Executive Summary Sheet
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    total_emps = db.query(Employee).count()
    high_risk_count = db.query(Employee).filter(Employee.risk_score >= 60).count()
    crit_risk_count = db.query(Employee).filter(Employee.risk_score >= 80).count()
    total_alerts = db.query(Alert).count()
    open_alerts = db.query(Alert).filter(Alert.status == "OPEN").count()
    total_incidents = db.query(Incident).count()
    total_telemetry = mongo_service.query_activity_logs(limit=1)["total"]

    ws_summary.append(["ITBIS — Insider Threat Behavioral Intelligence System"])
    ws_summary.append(["Executive SOC & Enterprise Threat Intelligence Report"])
    ws_summary.append([f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"])
    ws_summary.append([])

    ws_summary["A1"].font = title_font
    ws_summary["A2"].font = Font(name="Arial", size=11, italic=True, color="555555")
    ws_summary["A3"].font = Font(name="Arial", size=9, color="777777")

    summary_headers = ["Metric / Indicator", "Current Value", "Operational Status"]
    ws_summary.append(summary_headers)
    style_header_row(ws_summary, 5)

    kpi_rows = [
        ("Total Monitored Employees", total_emps, "ACTIVE"),
        ("High-Risk Subjects (Risk Score >= 60)", high_risk_count, "ATTENTION REQUIRED" if high_risk_count > 0 else "NORMAL"),
        ("Critical-Risk Subjects (Risk Score >= 80)", crit_risk_count, "CONTAINMENT RECOMMENDED" if crit_risk_count > 0 else "NOMINAL"),
        ("Total Security Alerts (All-Time)", total_alerts, "RECORDED"),
        ("Active Open Alerts", open_alerts, "INVESTIGATION QUEUED" if open_alerts > 0 else "CLEARED"),
        ("Total Forensic Incident Tickets", total_incidents, "MANAGED"),
        ("Total Activity Telemetry Events Ingested", total_telemetry, "LOGGED"),
        ("Scoring Model", "5-Factor Weighted Formula (35% BA, 25% PM, 20% DAV, 10% APD, 10% HSE)", "100% BALANCED"),
    ]

    for label, val, stat in kpi_rows:
        ws_summary.append([label, val, stat])

    for row in ws_summary.iter_rows(min_row=6, max_row=5+len(kpi_rows), min_col=1, max_col=3):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    auto_fit_columns(ws_summary)

    # 2. Employees Sheet
    ws_emp = wb.create_sheet(title="Employees")
    ws_emp.views.sheetView[0].showGridLines = True
    emp_headers = [
        "Employee ID", "Full Name", "Email Address", "Department",
        "Designation", "Employment Type", "Work Location", "Risk Score",
        "Risk Level", "Account Status", "Hardware Device", "Access Privileges"
    ]
    ws_emp.append(emp_headers)
    style_header_row(ws_emp, 1)

    employees = db.query(Employee).order_by(Employee.risk_score.desc(), Employee.id.asc()).all()
    for emp in employees:
        ws_emp.append([
            emp.employee_id, emp.name, emp.email, emp.department or "",
            emp.designation or "", emp.employment_type or "Full-time",
            emp.work_location or "HQ", emp.risk_score, emp.risk_level,
            emp.status, emp.device_info or "N/A", emp.access_privileges or "STANDARD"
        ])

    for row in ws_emp.iter_rows(min_row=2, max_row=1+len(employees), min_col=1, max_col=len(emp_headers)):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    auto_fit_columns(ws_emp)

    # 3. Risk Scores History Sheet
    ws_risk = wb.create_sheet(title="Risk Scores")
    ws_risk.views.sheetView[0].showGridLines = True
    risk_headers = ["ID", "Employee ID", "Employee Name", "Department", "Risk Score", "Risk Level", "Calculation Date"]
    ws_risk.append(risk_headers)
    style_header_row(ws_risk, 1)

    histories = db.query(RiskScoreHistory).order_by(RiskScoreHistory.calculated_at.desc()).limit(200).all()
    for h in histories:
        emp_name = h.employee.name if h.employee else "Unknown"
        emp_id = h.employee.employee_id if h.employee else f"EMP-{h.employee_id}"
        emp_dept = h.employee.department if h.employee else "N/A"
        date_str = h.calculated_at.strftime("%Y-%m-%d %H:%M:%S") if h.calculated_at else ""
        ws_risk.append([h.id, emp_id, emp_name, emp_dept, h.score, h.risk_level, date_str])

    for row in ws_risk.iter_rows(min_row=2, max_row=1+len(histories), min_col=1, max_col=len(risk_headers)):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    auto_fit_columns(ws_risk)

    # 4. Alerts Sheet
    ws_alerts = wb.create_sheet(title="Alerts")
    ws_alerts.views.sheetView[0].showGridLines = True
    alert_headers = ["Alert ID", "Employee ID", "Employee Name", "Title", "Category", "Severity", "Risk Score", "Status", "Created At"]
    ws_alerts.append(alert_headers)
    style_header_row(ws_alerts, 1)

    alerts = db.query(Alert).order_by(Alert.created_at.desc()).all()
    for a in alerts:
        emp_name = a.employee.name if a.employee else "Unknown"
        emp_id = a.employee.employee_id if a.employee else f"EMP-{a.employee_id}"
        date_str = a.created_at.strftime("%Y-%m-%d %H:%M:%S") if a.created_at else ""
        ws_alerts.append([a.id, emp_id, emp_name, a.title, a.category, a.severity, a.risk_score, a.status, date_str])

    for row in ws_alerts.iter_rows(min_row=2, max_row=1+len(alerts), min_col=1, max_col=len(alert_headers)):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    auto_fit_columns(ws_alerts)

    # 5. Incidents Sheet
    ws_inc = wb.create_sheet(title="Incidents")
    ws_inc.views.sheetView[0].showGridLines = True
    inc_headers = ["Incident ID", "Employee ID", "Subject Name", "Title", "Severity", "Status", "Assigned Analyst", "Created At", "Resolved At"]
    ws_inc.append(inc_headers)
    style_header_row(ws_inc, 1)

    incidents = db.query(Incident).order_by(Incident.created_at.desc()).all()
    for inc in incidents:
        emp_name = inc.employee.name if inc.employee else "Unknown"
        emp_id = inc.employee.employee_id if inc.employee else f"EMP-{inc.employee_id}"
        c_date = inc.created_at.strftime("%Y-%m-%d %H:%M:%S") if inc.created_at else ""
        r_date = inc.resolved_at.strftime("%Y-%m-%d %H:%M:%S") if inc.resolved_at else "Active"
        ws_inc.append([inc.id, emp_id, emp_name, inc.title, inc.severity, inc.status, inc.assigned_to or "SOC Lead", c_date, r_date])

    for row in ws_inc.iter_rows(min_row=2, max_row=1+len(incidents), min_col=1, max_col=len(inc_headers)):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    auto_fit_columns(ws_inc)

    # 6. Activity Logs Sheet
    ws_act = wb.create_sheet(title="Activity Logs")
    ws_act.views.sheetView[0].showGridLines = True
    act_headers = ["Event Type", "Employee ID", "Action", "Resource Target", "Device Node", "IP Address", "Timestamp", "Risk Weight", "Suspicious"]
    ws_act.append(act_headers)
    style_header_row(ws_act, 1)

    recent_acts = mongo_service.query_activity_logs(limit=100)["items"]
    for act in recent_acts:
        ws_act.append([
            act.get("event_type", ""),
            act.get("employee_id", ""),
            act.get("action", ""),
            act.get("resource", ""),
            act.get("device", ""),
            act.get("ip_address", ""),
            act.get("timestamp", ""),
            act.get("risk_weight", 0),
            "YES" if act.get("is_suspicious") else "NO"
        ])

    for row in ws_act.iter_rows(min_row=2, max_row=1+len(recent_acts), min_col=1, max_col=len(act_headers)):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    auto_fit_columns(ws_act)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    xlsx_bytes = output.getvalue()

    filename = f"ITBIS_Enterprise_Risk_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=xlsx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/export-pdf")
def export_reports_pdf(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Generates a genuine binary PDF security & compliance report using ReportLab."""
    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    # Custom Cyber SOC Styles
    title_style = ParagraphStyle(
        "ITBISTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0B1626"),
        spaceAfter=4,
    )

    subtitle_style = ParagraphStyle(
        "ITBISSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#00A8FF"),
        spaceAfter=12,
    )

    section_style = ParagraphStyle(
        "ITBISSection",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0B1626"),
        spaceBefore=12,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        "ITBISBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#222222"),
    )

    meta_style = ParagraphStyle(
        "ITBISMeta",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#555555"),
    )

    story = []

    # Title & Metadata Banner
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    story.append(Paragraph("ITBIS — INSIDER THREAT BEHAVIORAL INTELLIGENCE SYSTEM", title_style))
    story.append(Paragraph("EXECUTIVE THREAT & SOC COMPLIANCE DOSSIER // CONFIDENTIAL", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#00A8FF"), spaceAfter=10))

    meta_data = [
        [
            Paragraph(f"<b>Generated At:</b> {now_str}", meta_style),
            Paragraph("<b>Classification:</b> RESTRICTED // SOC TIER-3", meta_style),
        ],
        [
            Paragraph("<b>Compliance:</b> SOC 2 Type II / NIST 800-53 / ISO 27001", meta_style),
            Paragraph("<b>Scoring Model:</b> 5-Factor Weighted Behavioral Bayesian Sum", meta_style),
        ]
    ]
    meta_table = Table(meta_data, colWidths=[270, 270])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F0F4F8")),
        ('PADDING', (0,0), (-1,-1), 4),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 10))

    # 1. Executive Summary Table
    story.append(Paragraph("1. Executive Telemetry & Threat Posture Summary", section_style))

    total_emps = db.query(Employee).count()
    high_risk_count = db.query(Employee).filter(Employee.risk_score >= 60).count()
    crit_risk_count = db.query(Employee).filter(Employee.risk_score >= 80).count()
    total_alerts = db.query(Alert).count()
    open_alerts = db.query(Alert).filter(Alert.status == "OPEN").count()
    total_incidents = db.query(Incident).count()

    summary_data = [
        [Paragraph("<b>Key Security Metric</b>", meta_style), Paragraph("<b>Value</b>", meta_style), Paragraph("<b>Posture Assessment</b>", meta_style)],
        [Paragraph("Monitored Workforce", body_style), Paragraph(str(total_emps), body_style), Paragraph("ACTIVE MONITORING", body_style)],
        [Paragraph("High-Risk Subjects (Score &gt;= 60)", body_style), Paragraph(str(high_risk_count), body_style), Paragraph("UNDER ACTIVE REVIEW" if high_risk_count > 0 else "NOMINAL", body_style)],
        [Paragraph("Critical-Risk Subjects (Score &gt;= 80)", body_style), Paragraph(str(crit_risk_count), body_style), Paragraph("CONTAINMENT REQUIRED" if crit_risk_count > 0 else "NONE DETECTED", body_style)],
        [Paragraph("Total Confirmed Alerts", body_style), Paragraph(str(total_alerts), body_style), Paragraph(f"{open_alerts} Open in SOC Queue", body_style)],
        [Paragraph("Total Incidents Opened", body_style), Paragraph(str(total_incidents), body_style), Paragraph("Lifecycle Managed", body_style)],
    ]
    summary_table = Table(summary_data, colWidths=[200, 100, 240])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0B1626")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor("#00E5FF")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D0D7DE")),
        ('PADDING', (0,0), (-1,-1), 4),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 10))

    # 2. 5-Factor Risk Scoring Model Explanation
    story.append(Paragraph("2. 5-Factor Mathematical Risk Weighting Architecture", section_style))
    formula_data = [
        [Paragraph("<b>Factor Vector</b>", meta_style), Paragraph("<b>Weight</b>", meta_style), Paragraph("<b>Analysis Focus</b>", meta_style)],
        [Paragraph("1. Behavioral Anomalies (BA)", body_style), Paragraph("35%", body_style), Paragraph("Off-hours bursts, foreign devices, volume surges", body_style)],
        [Paragraph("2. Privilege Misuse Indicators (PM)", body_style), Paragraph("25%", body_style), Paragraph("Unauthorized sudo, root, Active Directory admin elevation", body_style)],
        [Paragraph("3. Data Access Violations (DAV)", body_style), Paragraph("20%", body_style), Paragraph("NTDS dumps, confidential SharePoint downloads, USB exfiltration", body_style)],
        [Paragraph("4. Access Pattern Deviations (APD)", body_style), Paragraph("10%", body_style), Paragraph("Untrusted public subnets, foreign VPN gateways", body_style)],
        [Paragraph("5. Historical Security Events (HSE)", body_style), Paragraph("10%", body_style), Paragraph("Prior confirmed security alerts and repeat incidents", body_style)],
    ]
    formula_table = Table(formula_data, colWidths=[180, 60, 300])
    formula_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0B1626")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor("#00E5FF")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D0D7DE")),
        ('PADDING', (0,0), (-1,-1), 3),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(formula_table)
    story.append(Spacer(1, 10))

    # 3. High-Risk Subject Dossiers Table
    story.append(Paragraph("3. Monitored Subjects Risk Overview", section_style))
    high_emps = db.query(Employee).order_by(Employee.risk_score.desc()).limit(8).all()
    emp_table_data = [
        [Paragraph("<b>ID</b>", meta_style), Paragraph("<b>Name</b>", meta_style), Paragraph("<b>Department</b>", meta_style), Paragraph("<b>Score</b>", meta_style), Paragraph("<b>Tier</b>", meta_style), Paragraph("<b>Status</b>", meta_style)]
    ]
    for emp in high_emps:
        emp_table_data.append([
            Paragraph(emp.employee_id, body_style),
            Paragraph(emp.name, body_style),
            Paragraph(emp.department or "N/A", body_style),
            Paragraph(str(emp.risk_score), body_style),
            Paragraph(emp.risk_level, body_style),
            Paragraph(emp.status, body_style),
        ])

    emp_table = Table(emp_table_data, colWidths=[65, 140, 115, 60, 80, 80])
    emp_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0B1626")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor("#00E5FF")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D0D7DE")),
        ('PADDING', (0,0), (-1,-1), 3),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(emp_table)
    story.append(Spacer(1, 10))

    # 4. Recent High Severity Alerts
    story.append(Paragraph("4. Critical & High Security Alerts", section_style))
    top_alerts = db.query(Alert).filter(Alert.severity.in_(["CRITICAL", "HIGH"])).order_by(Alert.created_at.desc()).limit(6).all()
    if top_alerts:
        alert_table_data = [
            [Paragraph("<b>Alert ID</b>", meta_style), Paragraph("<b>Subject</b>", meta_style), Paragraph("<b>Category</b>", meta_style), Paragraph("<b>Severity</b>", meta_style), Paragraph("<b>Status</b>", meta_style)]
        ]
        for a in top_alerts:
            emp_name = a.employee.name if a.employee else "Unknown"
            alert_table_data.append([
                Paragraph(f"ALT-{a.id:03d}", body_style),
                Paragraph(emp_name, body_style),
                Paragraph(a.category, body_style),
                Paragraph(a.severity, body_style),
                Paragraph(a.status, body_style),
            ])
        alert_table = Table(alert_table_data, colWidths=[65, 130, 165, 90, 90])
        alert_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0B1626")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor("#00E5FF")),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D0D7DE")),
            ('PADDING', (0,0), (-1,-1), 3),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        story.append(alert_table)
    else:
        story.append(Paragraph("No critical or high severity alerts registered in the current evaluation window.", body_style))

    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CCCCCC"), spaceAfter=6))
    story.append(Paragraph("<b>CONFIDENTIAL DOCUMENT</b> — Generated automatically by ITBIS Security Platform. Unauthorized redistribution is strictly prohibited.", meta_style))

    doc.build(story)
    pdf_buffer.seek(0)
    pdf_bytes = pdf_buffer.getvalue()

    filename = f"ITBIS_SOC_Executive_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/catalog")
def get_reports_catalog(
    current_user: dict = Depends(require_role("admin", "security_analyst", "security_manager")),
):
    """
    Catalog of all reporting endpoints, supported formats, and access roles (Day 33-34).
    """
    return [
        {
            "id": "employee_risk_dossier",
            "name": "Employee Forensic Risk Dossier",
            "format": "PDF",
            "endpoint": "/reports/{employee_id}/pdf",
            "description": "Comprehensive employee risk profile, 5-factor breakdown, and recent forensic anomalies.",
            "roles": ["admin", "security_analyst", "security_manager"],
        },
        {
            "id": "compliance_sla_report",
            "name": "SOC Compliance & SLA Performance Audit",
            "format": "PDF",
            "endpoint": "/reports/compliance/pdf",
            "description": "Executive compliance audit covering MTTR, MTTI, incident resolution rate, and SLA adherence.",
            "roles": ["admin", "security_manager"],
        },
        {
            "id": "export_risk_scores",
            "name": "Enterprise Risk Scores Export",
            "format": "XLSX",
            "endpoint": "/reports/export/risk-scores",
            "description": "Bulk export of monitored employee risk scores, sorted descending with department filtering.",
            "roles": ["admin", "security_analyst", "security_manager"],
        },
        {
            "id": "export_incidents",
            "name": "Forensic Incidents Export",
            "format": "XLSX",
            "endpoint": "/reports/export/incidents",
            "description": "Full lifecycle export of security incidents with resolution timestamps and status filtering.",
            "roles": ["admin", "security_analyst", "security_manager"],
        },
        {
            "id": "executive_summary",
            "name": "Executive Threat Intelligence Summary",
            "format": "JSON",
            "endpoint": "/reports/summary",
            "description": "Real-time threat landscape, incident ratios, and top threat category breakdown.",
            "roles": ["admin", "security_analyst", "security_manager"],
        },
        {
            "id": "soc_executive_pdf",
            "name": "SOC Executive Multi-Page Dossier",
            "format": "PDF",
            "endpoint": "/reports/export-pdf",
            "description": "Enterprise-level executive PDF report with scoring models, alerts, and subject lists.",
            "roles": ["admin", "security_analyst", "security_manager"],
        },
    ]


@router.get("/compliance/pdf")
def export_compliance_pdf(
    db: Session = Depends(get_db),
    user: dict = Depends(require_role("admin", "security_manager")),
):
    """
    Generates an in-memory compliance & SLA audit PDF report using ReportLab (Day 33-34).
    Restricted to Admins and Security Managers (Analysts get 403 Forbidden).
    """
    metrics = get_compliance_metrics(db)

    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "CompTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0B1626"),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "CompSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#00A8FF"),
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        "CompSection",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0B1626"),
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "CompBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#222222"),
    )
    meta_style = ParagraphStyle(
        "CompMeta",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#555555"),
    )

    story = []
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    story.append(Paragraph("ITBIS — SOC COMPLIANCE & SLA PERFORMANCE REPORT", title_style))
    story.append(Paragraph("CONFIDENTIAL // EXECUTIVE AUDIT & GOVERNANCE RECORD", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#00A8FF"), spaceAfter=10))

    meta_data = [
        [
            Paragraph(f"<b>Generated At:</b> {now_str}", meta_style),
            Paragraph(f"<b>Compliance Status:</b> {metrics.get('compliance_status', 'COMPLIANT')}", meta_style),
        ],
        [
            Paragraph("<b>Target Frameworks:</b> SOC 2 Type II / ISO 27001 / NIST CSF", meta_style),
            Paragraph(f"<b>Audited By:</b> {user.get('sub', 'System')}", meta_style),
        ]
    ]
    meta_table = Table(meta_data, colWidths=[270, 270])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F0F4F8")),
        ('PADDING', (0,0), (-1,-1), 4),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 10))

    story.append(Paragraph("1. SLA Response & Resolution Performance Metrics", section_style))

    kpi_data = [
        [Paragraph("<b>Metric Indicator</b>", meta_style), Paragraph("<b>Observed Value</b>", meta_style), Paragraph("<b>Target Threshold</b>", meta_style), Paragraph("<b>Status</b>", meta_style)],
        [Paragraph("Mean Time to Investigate (MTTI)", body_style), Paragraph(f"{metrics['mtti_hours']:.2f} hours", body_style), Paragraph("&lt; 4.0 hours", body_style), Paragraph("PASS" if metrics['mtti_hours'] <= 4.0 else "REVIEW", body_style)],
        [Paragraph("Mean Time to Resolve (MTTR)", body_style), Paragraph(f"{metrics['mttr_hours']:.2f} hours", body_style), Paragraph("&lt; 24.0 hours", body_style), Paragraph("PASS" if metrics['mttr_hours'] <= 24.0 else "REVIEW", body_style)],
        [Paragraph("Total Incidents Opened", body_style), Paragraph(str(metrics['total_incidents']), body_style), Paragraph("N/A", body_style), Paragraph("TRACKED", body_style)],
        [Paragraph("Resolved Incidents", body_style), Paragraph(str(metrics['resolved_incidents']), body_style), Paragraph("N/A", body_style), Paragraph("TRACKED", body_style)],
        [Paragraph("Incident Resolution Rate", body_style), Paragraph(f"{metrics['resolution_rate_percent']:.1f}%", body_style), Paragraph("&gt;= 75.0%", body_style), Paragraph("PASS" if metrics['resolution_rate_percent'] >= 75.0 else "ACTION REQUIRED", body_style)],
        [Paragraph("Active Critical Incidents", body_style), Paragraph(str(metrics['open_critical_incidents']), body_style), Paragraph("0", body_style), Paragraph("CLEAR" if metrics['open_critical_incidents'] == 0 else "CRITICAL ATTENTION", body_style)],
    ]

    kpi_table = Table(kpi_data, colWidths=[180, 100, 110, 150])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0B1626")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor("#00E5FF")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D0D7DE")),
        ('PADDING', (0,0), (-1,-1), 4),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 14))

    story.append(Paragraph("2. Compliance Governance Summary", section_style))
    gov_text = (
        f"This report certifies that the ITBIS platform evaluated {metrics['total_incidents']} incident cases. "
        f"The incident resolution rate is currently {metrics['resolution_rate_percent']:.1f}%, with Mean Time to Investigate (MTTI) "
        f"at {metrics['mtti_hours']:.2f} hours and Mean Time to Resolve (MTTR) at {metrics['mttr_hours']:.2f} hours. "
        f"The organization's posture is categorized as <b>{metrics['compliance_status']}</b> under continuous behavioral surveillance."
    )
    story.append(Paragraph(gov_text, body_style))
    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CCCCCC"), spaceAfter=6))
    story.append(Paragraph("<b>CONFIDENTIAL COMPLIANCE AUDIT</b> — Generated by ITBIS Automated Compliance Engine.", meta_style))

    doc.build(story)
    pdf_buffer.seek(0)
    pdf_bytes = pdf_buffer.getvalue()

    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="compliance_sla_report.pdf"'}
    )


@router.get("/export/risk-scores")
def export_risk_scores_excel(
    department: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    user: dict = Depends(require_role("admin", "security_analyst", "security_manager")),
):
    """
    Exports monitored employees' risk scores to an Excel spreadsheet (Day 33-34).
    Sorted descending by risk score, with optional department filtering.
    """
    query = db.query(Employee)
    if department and department.strip().upper() != "ALL":
        query = query.filter(Employee.department == department.strip())

    employees = query.order_by(Employee.risk_score.desc(), Employee.id.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk Scores"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="0B1626", end_color="0B1626", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="00E5FF")
    data_font = Font(name="Arial", size=9)
    border_thin = Border(
        left=Side(style="thin", color="E0E0E0"), right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"), bottom=Side(style="thin", color="E0E0E0")
    )

    headers = ["Employee ID", "Name", "Department", "Designation", "Risk Score", "Risk Level", "Status"]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for emp in employees:
        ws.append([
            emp.employee_id,
            emp.name,
            emp.department or "N/A",
            emp.designation or "N/A",
            emp.risk_score or 0,
            emp.risk_level or "LOW",
            emp.status or "ACTIVE",
        ])

    for row in ws.iter_rows(min_row=2, max_row=1 + len(employees), min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="employee_risk_scores.xlsx"'},
    )


@router.get("/export/incidents")
def export_incidents_excel(
    status_filter: Optional[str] = Query(None, alias="status"),
    db: Session = Depends(get_db),
    user: dict = Depends(require_role("admin", "security_analyst", "security_manager")),
):
    """
    Exports incident investigations to an Excel spreadsheet (Day 33-34).
    Unresolved incidents explicitly display 'Not resolved'.
    """
    query = db.query(Incident)
    if status_filter and status_filter.strip().upper() != "ALL":
        query = query.filter(Incident.status.ilike(status_filter.strip()))

    incidents = query.order_by(Incident.created_at.desc(), Incident.id.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Incidents"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="0B1626", end_color="0B1626", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="00E5FF")
    data_font = Font(name="Arial", size=9)
    border_thin = Border(
        left=Side(style="thin", color="E0E0E0"), right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"), bottom=Side(style="thin", color="E0E0E0")
    )

    headers = [
        "Incident ID", "Employee ID", "Employee Name", "Title",
        "Severity", "Status", "Assigned To", "Created At", "Resolved At"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for inc in incidents:
        emp_name = inc.employee.name if inc.employee else "Unknown"
        emp_id = inc.employee.employee_id if inc.employee else f"EMP-{inc.employee_id}"
        created_str = inc.created_at.strftime("%Y-%m-%d %H:%M:%S") if inc.created_at else ""
        resolved_str = inc.resolved_at.strftime("%Y-%m-%d %H:%M:%S") if inc.resolved_at else "Not resolved"

        ws.append([
            inc.id,
            emp_id,
            emp_name,
            inc.title or "N/A",
            inc.severity or "HIGH",
            inc.status or "OPEN",
            str(inc.assigned_to) if inc.assigned_to else "Unassigned",
            created_str,
            resolved_str,
        ])

    for row in ws.iter_rows(min_row=2, max_row=1 + len(incidents), min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="incidents_export.xlsx"'},
    )


@router.get("/{employee_id}/pdf")
def export_employee_risk_pdf(
    employee_id: str,
    db: Session = Depends(get_db),
    user: dict = Depends(require_role("admin", "security_analyst", "security_manager")),
):
    """
    Generates an in-memory forensic PDF risk dossier for a specific employee (Day 33-34).
    Contains risk score, 5-factor breakdown table, and first 10 anomalies.
    Returns 404 if employee does not exist.
    """
    employee = db.query(Employee).filter(
        or_(
            Employee.employee_id == employee_id,
            Employee.id == int(employee_id) if employee_id.isdigit() else False,
        )
    ).first()

    if not employee:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Employee {employee_id} not found"
        )

    risk_data = calculate_risk_score(employee.employee_id, db)
    sub_scores = risk_data.get("sub_scores", {})
    overall_score = risk_data.get("final_score", employee.risk_score or 0)
    risk_level = risk_data.get("risk_level", employee.risk_level or "Low")

    # Get first 10 anomalies for this employee
    all_anomalies = mongo_service.get_anomalies(limit=5000)
    emp_anomalies = [
        a for a in all_anomalies
        if str(a.get("employee_id") or "").strip() == employee.employee_id
        or str(a.get("employee_id") or "").strip() == str(employee.id)
    ][:10]

    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "EmpTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0B1626"),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "EmpSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#00A8FF"),
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        "EmpSection",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0B1626"),
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "EmpBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#222222"),
    )
    meta_style = ParagraphStyle(
        "EmpMeta",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#555555"),
    )

    story = []
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    story.append(Paragraph(f"ITBIS — EMPLOYEE FORENSIC DOSSIER: {employee.employee_id}", title_style))
    story.append(Paragraph("INSIDER THREAT BEHAVIORAL INTELLIGENCE RECORD // RESTRICTED", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#00A8FF"), spaceAfter=10))

    # Employee profile metadata
    meta_data = [
        [
            Paragraph(f"<b>Employee Name:</b> {employee.name}", meta_style),
            Paragraph(f"<b>Employee ID:</b> {employee.employee_id}", meta_style),
        ],
        [
            Paragraph(f"<b>Department:</b> {employee.department or 'N/A'}", meta_style),
            Paragraph(f"<b>Designation:</b> {employee.designation or 'N/A'}", meta_style),
        ],
        [
            Paragraph(f"<b>Overall Risk Score:</b> <font color='red'><b>{overall_score} / 100</b></font>", meta_style),
            Paragraph(f"<b>Threat Category / Tier:</b> <b>{risk_level}</b>", meta_style),
        ],
        [
            Paragraph(f"<b>Generated At:</b> {now_str}", meta_style),
            Paragraph(f"<b>Generated By:</b> {user.get('sub', 'Analyst')}", meta_style),
        ]
    ]
    meta_table = Table(meta_data, colWidths=[270, 270])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F0F4F8")),
        ('PADDING', (0,0), (-1,-1), 4),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 10))

    # 1. 5-Factor Weighted Sub-Scores Table
    story.append(Paragraph("1. 5-Factor Weighted Risk Component Breakdown", section_style))

    factors = [
        ("Behavioral Anomalies (BA)", "35%", sub_scores.get("behavioral_anomalies", 0.0)),
        ("Privilege Misuse Indicators (PM)", "25%", sub_scores.get("privilege_misuse", 0.0)),
        ("Data Access Violations (DAV)", "20%", sub_scores.get("data_access_violations", 0.0)),
        ("Access Pattern Deviations (APD)", "10%", sub_scores.get("access_pattern_deviations", 0.0)),
        ("Historical Security Events (HSE)", "10%", sub_scores.get("historical_events", 0.0)),
    ]

    breakdown_data = [
        [
            Paragraph("<b>Risk Factor Vector</b>", meta_style),
            Paragraph("<b>Model Weight</b>", meta_style),
            Paragraph("<b>Raw Sub-Score (0-100)</b>", meta_style),
            Paragraph("<b>Weighted Points</b>", meta_style),
        ]
    ]
    for name, weight, raw_val in factors:
        wt_num = float(weight.replace("%", "")) / 100.0
        weighted_pts = round(raw_val * wt_num, 2)
        breakdown_data.append([
            Paragraph(name, body_style),
            Paragraph(weight, body_style),
            Paragraph(f"{raw_val:.1f}", body_style),
            Paragraph(f"{weighted_pts:.2f}", body_style),
        ])

    breakdown_data.append([
        Paragraph("<b>Total Composite Risk Score</b>", meta_style),
        Paragraph("<b>100%</b>", meta_style),
        Paragraph("-", meta_style),
        Paragraph(f"<b>{overall_score}</b>", meta_style),
    ])

    breakdown_table = Table(breakdown_data, colWidths=[240, 80, 110, 110])
    breakdown_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0B1626")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor("#00E5FF")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D0D7DE")),
        ('PADDING', (0,0), (-1,-1), 4),
        ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor("#F8FAFC")]),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor("#E2E8F0")),
    ]))
    story.append(breakdown_table)
    story.append(Spacer(1, 12))

    # 2. Recent Forensic Anomalies Table (first 10)
    story.append(Paragraph("2. Recent Forensic Anomalies (First 10 Observed)", section_style))
    if emp_anomalies:
        anom_data = [
            [
                Paragraph("<b>Date / Timestamp</b>", meta_style),
                Paragraph("<b>Anomaly Type</b>", meta_style),
                Paragraph("<b>Severity</b>", meta_style),
                Paragraph("<b>Details / Description</b>", meta_style),
            ]
        ]
        for a in emp_anomalies:
            ts = str(a.get("detected_at") or a.get("detection_timestamp") or a.get("timestamp") or "N/A")[:19]
            atype = str(a.get("anomaly_type") or "ANOMALY").replace("_", " ")
            sev = str(a.get("severity") or "MEDIUM").upper()
            details_dict = a.get("details") if isinstance(a.get("details"), dict) else {}
            desc = a.get("description") or details_dict.get("description") or details_dict.get("reason") or "Behavioral deviation recorded."
            anom_data.append([
                Paragraph(ts, body_style),
                Paragraph(atype, body_style),
                Paragraph(sev, body_style),
                Paragraph(str(desc)[:75], body_style),
            ])

        anom_table = Table(anom_data, colWidths=[110, 130, 80, 220])
        anom_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0B1626")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor("#00E5FF")),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D0D7DE")),
            ('PADDING', (0,0), (-1,-1), 3),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        story.append(anom_table)
    else:
        story.append(Paragraph("No telemetry anomalies currently registered for this employee.", body_style))

    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CCCCCC"), spaceAfter=6))
    story.append(Paragraph("<b>CONFIDENTIAL DOSSIER</b> — Generated by ITBIS Forensic Engine. Unauthorized reproduction is strictly prohibited.", meta_style))

    doc.build(story)
    pdf_buffer.seek(0)
    pdf_bytes = pdf_buffer.getvalue()

    filename = f"employee_{employee.employee_id}_risk_report.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


